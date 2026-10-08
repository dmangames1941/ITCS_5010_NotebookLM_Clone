import os
import requests
from typing import List, Dict, Any
from bs4 import BeautifulSoup
from pypdf import PdfReader
from pptx import Presentation
from langchain_text_splitters import RecursiveCharacterTextSplitter

class DocumentIngestor:
    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 150):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", ""]
        )

    def extract_from_pdf(self, file_path: str) -> str:
        """Extracts text content from a PDF file using pypdf."""
        reader = PdfReader(file_path)
        text_blocks = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_blocks.append(page_text)
        return "\n\n".join(text_blocks)

    def extract_from_pptx(self, file_path: str) -> str:
        prs = Presentation(file_path)
        text_blocks = []
        for slide_idx, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    slide_text.append(shape.text.strip())
            if slide_text:
                text_blocks.append(f"--- Slide {slide_idx + 1} ---\n" + "\n".join(slide_text))
        return "\n\n".join(text_blocks)

    def extract_from_txt(self, file_path: str) -> str:
        """Reads plain text content from a file."""
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    def extract_from_url(self, url: str) -> str:
        """Scrapes text content from a public web URL using requests and BeautifulSoup."""
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64, x64) Application/RAG-Ingestor"
        }
        response = requests.get(url, headers=headers, timeout=15)
        response.raise_for_status()

        soup = BeautifulSoup(response.text, "html.parser")

        for element in soup(["script", "style", "nav", "footer", "header", "noscript"]):
            element.decompose()

        text = soup.get_text(separator="\n")

        lines = (line.strip() for line in text.splitlines())
        chunks = (phrase.strip() for line in lines for phrase in line.split(" "))
        return "\n".join(chunk for chunk in chunks if chunk)

    def process_source(self, source_input: str, source_type: str, notebook_id: str, source_name: str) -> List[Dict[str, Any]]:
        """
        Parses source material, chunks it, and attaches notebook/source metadata.

        Parameters:
            source_input: File path or Web URL string
            source_type: One of ['pdf', 'pptx', 'txt', 'url']
            notebook_id: ID of the owning notebook
            source_name: Display name/filename of the source
        
        Returns:
            List of chunk dictionaries ready for embedding and store
        """
        source_type = source_type.lower()

        if source_type == "pdf":
            raw_text = self.extract_from_pdf(source_input)
        elif source_type == "pptx":
            raw_text = self.extract_from_pptx(source_input)
        elif source_type == "txt":
            raw_text = self.extract_from_txt(source_input)
        elif source_type == "url":
            raw_text = self.extract_from_url(source_input)
        else:
            raise ValueError(f"Unsupported source type: {source_type}")

        if not raw_text.strip():
            raise ValueError(f"No extractable text found in source: {source_name}")

        chunks = self.text_splitter.split_text(raw_text)

        processed_chunks = []
        for idx, chunk in enumerate(chunks):
            processed_chunks.append({
                "text": chunk,
                "metadata": {
                    "notebook_id": notebook_id,
                    "source_name": source_name,
                    "source_type": source_type,
                    "chunk_index": idx,
                    "total_chunks": len(chunks)
                }
            })
        return processed_chunks